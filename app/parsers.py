from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

TEXT_EXTENSIONS = {
    ".txt", ".md", ".log", ".py", ".json", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".c", ".cc", ".cpp", ".h", ".hpp",
    ".rs", ".go", ".js", ".ts", ".tsx", ".java", ".sh"
}


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_csv(path: Path) -> str:
    rows: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as fh:
        for row in csv.reader(fh):
            rows.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(rows)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        return ""
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"[sheet:{sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            parts.append(" | ".join(values))
    return "\n".join(parts)


def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT_EXTENSIONS:
            return _read_text(path)
        if suffix == ".csv":
            return _read_csv(path)
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".docx":
            return _read_docx(path)
        if suffix in {".xlsx", ".xlsm"}:
            return _read_xlsx(path)
    except (OSError, PermissionError, ValueError, RuntimeError):
        return ""
    except Exception:
        return ""
    return ""


def iter_corpus_files(root: Path) -> Iterable[Path]:
    if not root.exists() or not root.is_dir():
        return []
    return (p for p in sorted(root.rglob("*")) if p.is_file())


def chunk_text(text: str, max_chars: int = 1800, overlap: int = 180) -> list[str]:
    compact = "\n".join(line.rstrip() for line in text.splitlines() if line.strip())
    if not compact:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(compact):
        end = min(len(compact), start + max_chars)
        chunks.append(compact[start:end])
        if end >= len(compact):
            break
        start = max(0, end - overlap)
    return chunks


def _csv_row_chunks(text: str) -> list[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return []
    header = lines[0]
    if len(lines) == 1:
        return [header]
    return [f"{header}\n{row}" for row in lines[1:]]


def _xlsx_row_chunks(text: str) -> list[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    chunks: list[str] = []
    sheet = ""
    header = ""
    for line in lines:
        if line.startswith("[sheet:"):
            sheet = line
            header = ""
            continue
        if not header:
            header = line
            continue
        prefix = "\n".join(part for part in (sheet, header) if part)
        chunks.append(f"{prefix}\n{line}" if prefix else line)
    if not chunks and lines:
        return lines
    return chunks


def document_chunks(path: Path, text: str) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _csv_row_chunks(text)
    if suffix in {".xlsx", ".xlsm"}:
        return _xlsx_row_chunks(text)
    return chunk_text(text)


def build_index(corpus: str) -> dict:
    root = Path(corpus).resolve()
    docs: list[dict] = []
    for path in iter_corpus_files(root):
        text = parse_file(path)
        if not text:
            continue
        try:
            rel = path.relative_to(root).as_posix()
        except ValueError:
            continue
        for idx, chunk in enumerate(document_chunks(path, text)):
            docs.append({"source": rel, "chunk_id": idx, "text": chunk})
    return {"corpus": str(root), "documents": docs}
